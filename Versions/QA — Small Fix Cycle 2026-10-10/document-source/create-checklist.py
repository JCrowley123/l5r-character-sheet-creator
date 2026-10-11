from pathlib import Path
import argparse
import sys
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--build', default='Awaiting deployment verification')
args = parser.parse_args()

pages = [
    ('iPhone preparation and print quality', [
        ('P1', 'Open the print preview',
         'Open a test character. In the sheet menu, tap Print. Inspect every page in the phone’s print preview.',
         'Text is dark and readable throughout. The five Ring cards have no element symbols or cream wash. No large olive patch obscures the page. Record the page count; different devices may paginate differently.'),
        ('P2', 'Check the whole character',
         'Compare the preview with the sheet. Check the character name, Clan, Family and School, then Skills, Advantages, Techniques, Combat and the last printed section. Include long entries if your character has them.',
         'The correct character and visible sheet content print without cut-off text or overlapping Identity labels. Clan, Family and School picker controls do not print misleading choices. Applied character details remain readable.'),
        ('P3', 'Cancel and return to play',
         'Cancel the print dialog. Return to the sheet and open Clan and School, then Rings and Traits.',
         'The sheet still works. Its on-screen Ring artwork and Clan, Family and School pickers are present. Cancelling has not changed character values.'),
    ]),
    ('iPhone PDF export', [
        ('P4', 'Export the selected saved character',
         'Prepare two saved test characters with clearly different names and Schools, for example Crab and Crane. Open one, then go to Characters. Open the other character’s menu and tap Export to PDF.',
         'The print preview contains the selected character, including the correct name and School. It never exports the previously open character. After cancelling print, the selected character remains open.'),
        ('P5', 'Keep current edits in the PDF',
         'On the currently open test character, change a harmless text field such as Background. Without switching characters, use that character’s Export to PDF action.',
         'The preview contains the newly entered text. Character values and edits remain intact after cancelling. Use only a test copy for this check.'),
        ('P6', 'Protect an unsaved character',
         'Optional, if you already have an unsaved test sheet with edits. From Characters, try Export to PDF on a different saved character. In the discard confirmation, choose Cancel. Create New Character in the wizard saves automatically, so it does not prepare this test.',
         'No print dialog opens. The unsaved character and its edits are still available. If you have no suitable unsaved sheet, leave Not run; do not delete saves or clear browser data to create one.'),
        ('P7', 'Save and reopen the PDF',
         'Export your test character again. Use the phone’s normal print preview and Share action to Save to Files. Open the saved PDF in Files and inspect every page. The exact preview controls may differ by iOS version.',
         'The PDF opens, is complete and readable, and belongs to the selected character. It has no Ring symbols, cream wash or olive patch. Record the filename and page count in Notes.'),
    ]),
    ('iPhone Skill Traits', [
        ('S1', 'Correct older saved rows',
         'On a test copy, set Intimidation to Willpower, Perform: Biwa to Awareness and Games: Go to Awareness to reproduce the old defaults. Record the ranks and spent XP, save, open another character, then reopen the test copy. An untouched pre-update backup may also be used.',
         'Intimidation becomes Awareness, Biwa becomes Agility and Go becomes Intelligence. Ranks and spent XP stay the same. This follows your ruling: old default values are corrected on opening, even if selected by hand. Other explicit Traits are tested in S4.'),
        ('S2', 'Add corrected Skills',
         'In Management, add Intimidation if it is absent. Add a blank Skill row and type Perform: Biwa. Add another and type Games: Go.',
         'Intimidation uses Awareness, Biwa uses Agility and Go uses Intelligence. Each row is usable and no duplicate or empty extra row appears unexpectedly.'),
        ('S3', 'Rename an automatic Trait',
         'Add a fresh blank Skill row and type Perform: Biwa. Once its Trait fills automatically, rename that row Games: Go without changing the Trait control yourself.',
         'The automatic Trait changes from Agility to Intelligence. The row’s rank is retained. Renaming it Perform: Song then gives Awareness.'),
        ('S4', 'Preserve a chosen Trait',
         'On a Perform: Biwa test row, choose Reflexes, rename it Games: Go, save and reopen. On another Biwa row, explicitly choose Awareness and rename it Games: Go. Check before saving, then save and reopen.',
         'Reflexes survives rename and reload. Explicit Awareness survives the rename in that session, then becomes Intelligence on reopening because it matches Go’s old default. Saved rows do not record how a Trait was chosen.'),
        ('S5', 'Use the corrected Trait in a roll',
         'Give your test character different Awareness and Agility values. Roll Perform: Biwa and inspect the roll preview. Cancel the roll after checking it.',
         'The preview uses Agility and the row’s Skill rank. It does not still use Awareness. Record the Trait values and displayed pool if the result is unexpected.'),
    ]),
    ('iPhone remaining checks and optional Windows checks', [
        ('S6', 'Read the Craft guidance',
         'Add Craft or a named Craft row such as Craft: Carpentry. Read the note below its name and open Skill Info. Also inspect Craft in Search.',
         'Its starting Trait remains Awareness. The guidance explains that the Trait varies by craft and the GM sets it, citing Core Rulebook p.143. It does not assign a guessed Trait to Carpentry.'),
        ('S7', 'Filter Skills in Search',
         'Open Search, choose Skills and inspect the Trait filter. Select Awareness and search for Intimidation. Open the result.',
         'Intimidation appears under Awareness and its detail agrees. Willpower is absent from the default Skills Trait options. This does not prevent a player choosing a different Trait on their own row.'),
        ('S8', 'Create a character through the wizard',
         'In Characters, tap Create New Character. Choose Crane and Doji Courtier. Finish the wizard and inspect the School’s Perform: Storytelling row.',
         'Perform: Storytelling is present with Awareness and its School-granted rank. The wizard finishes normally. Storytelling intentionally remains Awareness; S2 and S3 check the changed sub-skill Traits.'),
        ('W1', 'Repeat PDF export on Windows',
         'Optional. In your usual Windows browser, repeat P4 and save the result through the browser’s PDF printer. Inspect the first and last pages and the Ring and Identity areas.',
         'The selected character prints correctly, all pages are readable and no Ring symbols print. Record the browser and whether its background graphics option was enabled.'),
        ('W2', 'Repeat the Skill checks on Windows',
         'Optional. In the Windows browser, repeat S2, S3, S4 and S7 using a test character.',
         'Adding and renaming Skills, preserving an explicit Trait and Search filtering behave as described on the iPhone. Report individual failing IDs in Notes.'),
    ]),
]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(.62)
sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.68)
sec.right_margin = Inches(.68)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.04
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)

def para(text='', bold_prefix=None, keep=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    p.paragraph_format.keep_with_next = keep
    return p

def result_field(tid, sequence):
    p = para('Result: ', 'Result: ', True)
    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr')
    for key, val in [('alias', tid + ' result'), ('tag', tid + '_result'), ('id', str(2000 + sequence))]:
        element = OxmlElement('w:' + key)
        element.set(qn('w:val'), val)
        props.append(element)
    dropdown = OxmlElement('w:dropDownList')
    for choice in ('Not run', 'Pass', 'Fail'):
        item = OxmlElement('w:listItem')
        item.set(qn('w:displayText'), choice)
        item.set(qn('w:value'), choice)
        dropdown.append(item)
    props.append(dropdown)
    sdt.append(props)
    content = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'Not run'
    run.append(text)
    content.append(run)
    sdt.append(content)
    p._p.append(sdt)
    p.add_run('     (Pass / Fail / Not run)')
    return p

intro = [
    'Check the print repair, Export to PDF and the Skill Trait corrections on the live app. Use the iPhone first; Windows checks are optional. All results start as Not run.',
    'Live app: https://l5r-character-sheet-creator.pages.dev/',
    f'Build to test: {args.build}',
    'Test date: __________________    iPhone and iOS: __________________',
    'Open a fresh copy of the app after the update. Before changing anything, use Export JSON to back up your character. In its Characters menu, use Save As a copy and give the copy a test name. Keep the backup untouched.',
    'For each test, choose Pass, Fail or Not run in its Result field and add notes. If your editor does not support dropdowns, edit the text or use the companion Markdown checklist. Report failures by ID and what you saw. A skipped check stays Not run.',
]
doc.add_paragraph('L5R small fixes test checklist', 'Title')
for item in intro:
    para(item)

markdown = ['# L5R small fixes test checklist', '', *[p + '\n' for p in intro]]
sequence = 0
for page_i, (title, tests) in enumerate(pages):
    if page_i:
        doc.add_page_break()
    doc.add_paragraph(title, 'Heading 1')
    markdown.extend(['', '## ' + title, ''])
    for tid, title, steps, expected in tests:
        sequence += 1
        doc.add_paragraph(tid + '  ' + title, 'Heading 2')
        para('Do: ' + steps, 'Do: ', True)
        para('Pass when: ' + expected, 'Pass when: ', True)
        result_field(tid, sequence)
        para('Notes: __________________________________________________________', 'Notes: ')
        markdown.extend(['### ' + tid + ' ' + title, '', '**Do:** ' + steps, '', '**Pass when:** ' + expected, '', '**Result:** Not run (Pass / Fail / Not run)', '', '**Notes:**', ''])

footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('L5R test checklist  |  ')
run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
doc.core_properties.title = 'L5R small fixes test checklist'
doc.core_properties.subject = 'Owner checks for print repair PDF export and Skill Traits'
doc.core_properties.author = 'L5R character sheet project'
doc.core_properties.keywords = 'iPhone PDF Skill Traits checklist'
ROOT.mkdir(parents=True, exist_ok=True)
doc.save(ROOT / 'TEST-CHECKLIST.docx')
(ROOT / 'TEST-CHECKLIST.md').write_text('\n'.join(markdown), encoding='utf-8')
print('Wrote TEST-CHECKLIST.docx and TEST-CHECKLIST.md')
checked = Document(ROOT / 'TEST-CHECKLIST.docx')
all_text = '\n'.join(p.text for p in checked.paragraphs)
assert all_text.count('Result: ') == 17
assert all_text.count('Notes: ') == 17
assert sum(p.style.name == 'Heading 2' for p in checked.paragraphs) == 17
with zipfile.ZipFile(ROOT / 'TEST-CHECKLIST.docx') as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith('.xml'):
            ET.fromstring(archive.read(name))
    document_xml = ET.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    controls = document_xml.findall('.//w:sdt', ns)
    assert len(controls) == 17
    assert len({control.find('w:sdtPr/w:id', ns).get(qn('w:val')) for control in controls}) == 17
    for control in controls:
        values = [item.get(qn('w:value')) for item in control.findall('w:sdtPr/w:dropDownList/w:listItem', ns)]
        assert values == ['Not run', 'Pass', 'Fail']
        assert ''.join(control.find('w:sdtContent', ns).itertext()) == 'Not run'
print('Verified 17 tests, 17 Result dropdowns initialized Not run, 17 Notes fields and valid Word package XML.')
